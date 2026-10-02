"""C997 deck: strip the WSQ layer convert_deck.py leaves behind (cover, K/A codes, schedule, feedback)."""
import re
from pptx import Presentation

OUT = "courseware/Business Transformation with AI Agents (C997)-v1.0.pptx"
p = Presentation(OUT)


def put(shape, *lines):
    """Rewrite a shape paragraph-by-paragraph through runs[0] so fonts survive."""
    paras = shape.text_frame.paragraphs
    assert len(paras) >= len(lines), (shape.shape_id, len(paras), lines)
    for pa, text in zip(paras, lines):
        runs = pa.runs
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    for pa in paras[len(lines):]:
        pa._p.getparent().remove(pa._p)


def shape(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)


def slide_titled(title):
    hits = [s for s in p.slides for sh in s.shapes
            if sh.has_text_frame and sh.text_frame.text == title]
    assert len(hits) == 1, (title, len(hits))
    return hits[0]


def replace_para(old, new):
    n = 0
    for s in p.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for pa in sh.text_frame.paragraphs:
                if "".join(r.text for r in pa.runs) == old:
                    pa.runs[0].text = new
                    for r in pa.runs[1:]:
                        r.text = ""
                    n += 1
    assert n, old


# Cover
cover = p.slides[0]
put(shape(cover, 8), "NON-WSQ · 2 DAYS")
put(shape(cover, 9), "COURSE SLIDES  ·  NON-WSQ")
put(shape(cover, 12), "Course Code: C997",
    "Conducted by Tertiary Infotech Academy Pte Ltd  ·  UEN 201200696W")

# Trainer bio, ground rules, LMS slide
replace_para("WSQ courses on generative AI, business innovation, automation and data science.",
             "Courses on generative AI, business innovation, automation and data science.")
replace_para("75% attendance is required.", "Attend both days in full to get the most from the course.")
replace_para("My Courses  ›  Digital Transformation and Business Innovati",
             "My Courses  ›  Business Transformation with AI Agents")
replace_para("Keep them open — the assessment is open book.",
             "Keep them open — you will refer to them in every lab.")

# Schedule slide
sched = slide_titled("Lesson Plan — 2 Days, 8 hours/day")
replace_para("Lesson Plan — 2 Days, 8 hours/day", "Lesson Plan — 2 Days, 7.5 hours/day")
replace_para("–  Digital attendance (AM) · Introductions · Outcomes", "–  Introductions · Outcomes · Course outline")
replace_para("–  Revision · TRAQOM survey · Final Assessment (WA + PP)", "–  Course summary · Q&A · Course feedback")

# K/A competency codes: eyebrows "K1 · THE SHIFT", topic strap "(K1, K2, A1, A2)", recap badges
EYEBROW = re.compile(r"^(?:[KA]\d(?:,\s*)?)+\s*·\s*")
STRAP = re.compile(r"\s*\((?:[KA]\d(?:,\s*)?)+\)$")
BADGE = re.compile(r"^(?:[KA]\d(?:,\s*)?)+$")
for s in p.slides:
    badges = []
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for pa in sh.text_frame.paragraphs:
            t = "".join(r.text for r in pa.runs)
            new = STRAP.sub("", EYEBROW.sub("", t))
            if BADGE.match(t):
                badges.append(pa)
            elif new != t:
                pa.runs[0].text = new
                for r in pa.runs[1:]:
                    r.text = ""
    for k, pa in enumerate(badges, 1):          # recap cards: competency code -> card number
        pa.runs[0].text = str(k)
        for r in pa.runs[1:]:
            r.text = ""

# Follow-on courses: drop the WSQ programme prefix
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for pa in sh.text_frame.paragraphs:
                if pa.runs and pa.runs[0].text.startswith("WSQ - "):
                    pa.runs[0].text = pa.runs[0].text[len("WSQ - "):]

# Digital Attendance slide -> Course Feedback (rewritten in place; card design kept)
fb = slide_titled("Digital Attendance (Mandatory)")
for sid, text in [
    (4, "BEFORE YOU LEAVE"),
    (5, "Course Feedback"),
    (33, "Scan the feedback QR code"),
    (34, "Displayed by your trainer at the end of the course."),
    (38, "LIKED"),
    (39, "Tell us what worked well for you."),
    (43, "IMPROVE"),
    (44, "Tell us what we could do better."),
    (48, "APPLY"),
    (49, "Share how you will use AI agents at work."),
    (53, "2 MIN"),
    (54, "Your feedback shapes the next run of this course."),
]:
    put(shape(fb, sid), text)

# Move the feedback slide to sit just before Thank You!
ids = p.slides._sldIdLst
fb_el = next(el for el, s in zip(list(ids), p.slides) if s is fb)
ids.remove(fb_el)
ids.insert(len(ids) - 1, fb_el)

p.save(OUT)
print("fixed", OUT)
