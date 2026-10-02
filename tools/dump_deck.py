"""Dump shapes (id, paragraphs) of the given 0-based slides: python3 tools/dump_deck.py <pptx> 0 1 8"""
import sys
from pptx import Presentation
p = Presentation(sys.argv[1])
for i in map(int, sys.argv[2:]):
    s = p.slides[i]
    print(f"===== slide {i}")
    for sh in s.shapes:
        if sh.has_text_frame:
            paras = ["".join(r.text for r in pa.runs) for pa in sh.text_frame.paragraphs]
            print(f"  [{sh.shape_id}] {paras}")
        elif sh.has_table:
            for r in sh.table.rows:
                print(f"  [{sh.shape_id}] ROW {[c.text for c in r.cells]}")
        elif sh.shape_type == 13:
            print(f"  [{sh.shape_id}] PICTURE {sh.name}")
