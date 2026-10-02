"""List slide/shape/paragraph hits for a regex in the deck (incl. tables + notes)."""
import re, sys
from pptx import Presentation
p = Presentation(sys.argv[1]); rx = re.compile(sys.argv[2], re.I)
for i, s in enumerate(p.slides):
    for sh in s.shapes:
        if sh.has_text_frame:
            for k, pa in enumerate(sh.text_frame.paragraphs):
                t = "".join(r.text for r in pa.runs)
                if rx.search(t): print(f"{i:3} [{sh.shape_id}.{k}] {t}")
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    if rx.search(c.text): print(f"{i:3} [{sh.shape_id} TBL] {c.text}")
    if s.has_notes_slide and rx.search(s.notes_slide.notes_text_frame.text):
        print(f"{i:3} NOTES: {s.notes_slide.notes_text_frame.text[:200]}")
